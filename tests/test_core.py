import os
import queue
import shutil
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
DATA = Path(os.environ["TDK_DATA_DIR"])
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def _fresh_db(name):
    """Trỏ db sang thư mục trống riêng của test và khởi tạo."""
    import db as dbmod

    # Thư mục riêng cho từng test: kết nối của db không đóng ngay nên Windows không xóa được DB cũ.
    data = Path(tempfile.mkdtemp(prefix=f"{name}-", dir=DATA))
    dbmod.DB_PATH = data / "data.db"
    dbmod.init()
    return dbmod


def _make_scraper():
    """Scraper có hàng đợi và khóa thật nhưng không chạy luồng nền, không mở trình duyệt."""
    from scraper import Scraper

    scraper = Scraper.__new__(Scraper)
    scraper.jobs = queue.Queue()
    scraper.pending = {}
    scraper.current = None
    scraper.lock = threading.Lock()
    scraper._pw = None
    scraper._ctx = None
    scraper._ctx_visible = None
    scraper.RETRY_DELAYS = (0, 0)
    return scraper


def _wait_until(condition, timeout=10):
    deadline = timeout
    while deadline > 0:
        if condition():
            return True
        threading.Event().wait(0.02)
        deadline -= 0.02
    return condition()


def _run_worker(scraper, condition):
    """Chạy vòng _run thật (bỏ qua nhịp nghỉ 1,5 giây) cho đến khi điều kiện đúng."""
    with mock.patch("time.sleep", lambda seconds: None):
        threading.Thread(target=scraper._run, daemon=True).start()
        return _wait_until(lambda: condition() and scraper.current is None)


def _video(video_id, create_time=100, play=1, **extra):
    row = {
        "id": video_id, "description": "", "create_time": create_time, "duration": 5, "cover": "",
        "play": play, "digg": 0, "comment": 0, "share": 0, "collect": 0, "music": "", "is_ad": 0,
    }
    row.update(extra)
    return row


class ParseInputTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from scraper import parse_input

        cls.parse_input = staticmethod(parse_input)

    def test_username_url_and_junk(self):
        found, invalid = self.parse_input("@Demo.User, https://www.tiktok.com/@Other_1\nnope!")
        self.assertEqual(found, ["demo.user", "other_1"])
        self.assertEqual(invalid, ["nope!"])

    def test_dedupes_same_channel(self):
        found, invalid = self.parse_input("demo.user @demo.user")
        self.assertEqual(found, ["demo.user"])
        self.assertEqual(invalid, [])

    def test_wrapped_and_trailing_punctuation(self):
        found, invalid = self.parse_input('(@Demo.User) "@other_1" @third! (https://www.tiktok.com/@Fourth_4).')
        self.assertEqual(found, ["demo.user", "other_1", "third", "fourth_4"])
        self.assertEqual(invalid, [])

    def test_uppercase_url(self):
        found, invalid = self.parse_input("HTTPS://WWW.TIKTOK.COM/@Demo.User")
        self.assertEqual(found, ["demo.user"])
        self.assertEqual(invalid, [])

    def test_zero_width_and_fullwidth_at(self):
        found, invalid = self.parse_input("@demo​.user ＠Other_1 ​third")
        self.assertEqual(found, ["demo.user", "other_1", "third"])
        self.assertEqual(invalid, [])

    def test_dedupes_across_noisy_forms(self):
        text = 'demo.user (@Demo.User) "@demo.user" @demo.user! ＠demo​.user HTTPS://WWW.TIKTOK.COM/@DEMO.USER'
        self.assertEqual(self.parse_input(text), (["demo.user"], []))

    def test_invalid_stays_invalid(self):
        found, invalid = self.parse_input("nope! (x) @a$b")
        self.assertEqual(found, [])
        self.assertEqual(invalid, ["nope!", "(x)", "@a$b"])

    def test_empty(self):
        self.assertEqual(self.parse_input("  \n,"), ([], []))


class DbTest(unittest.TestCase):
    def test_tables_and_channel_roundtrip(self):
        import sqlite3

        db = _fresh_db("schema")
        added, existed = db.add_channels(["demo_user", "demo_user"])
        self.assertEqual(added, ["demo_user"])
        self.assertEqual(existed, ["demo_user"])
        db.save_channel_info("demo_user", {"followers": 10, "likes": 2, "video_count": 1})
        db.finish_scan("demo_user", [_video("v1", play=100)], full=True)
        db.save_channel_info("demo_user", {"followers": 12, "likes": 3})
        db.finish_scan("demo_user", [_video("v1", play=130)], full=False)
        row = next(item for item in db.list_channels() if item["username"] == "demo_user")
        self.assertEqual(row["total_views"], 130)
        self.assertEqual(row["d_followers"], 2)
        self.assertEqual(row["d_views"], 30)
        self.assertIsNotNone(row["prev_ts"])
        names = [
            r[0]
            for r in sqlite3.connect(db.DB_PATH).execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
            )
        ]
        self.assertEqual(names, ["alerts", "channels", "settings", "snapshots", "videos"])
        # Lần cập nhật sau của cùng video giữ lại số view trước đó để tính mức tăng.
        self.assertEqual(db.list_videos("demo_user")[0]["prev_play"], 100)

    def test_keys_are_case_insensitive_and_unique_id_keeps_case(self):
        db = _fresh_db("case")
        db.add_channels(["Demo.User"])
        self.assertTrue(db.channel_exists("demo.user"))
        self.assertEqual(db.get_channel("DEMO.user")["unique_id"], "Demo.User")
        self.assertEqual(db.add_channels(["demo.USER"]), ([], ["demo.USER"]))
        db.delete_channels(["DEMO.USER"])
        self.assertIsNone(db.get_channel("demo.user"))

    def test_channels_list_in_added_order(self):
        db = _fresh_db("order")
        db.add_channels(["first", "second", "third"])
        self.assertEqual([c["username"] for c in db.list_channels()], ["first", "second", "third"])

    def test_engagement_rate_is_percent_and_includes_saves(self):
        db = _fresh_db("er")
        db.add_channels(["demo"])
        db.finish_scan("demo", [_video("a", play=1000, digg=40, comment=5, share=3, collect=2)], full=False)
        row = db.list_channels()[0]
        self.assertEqual(row["er"], 5.0)
        self.assertEqual(row["avg_views"], 1000)
        self.assertEqual(row["top_views"], 1000)

    def test_channel_without_videos_has_no_aggregates(self):
        db = _fresh_db("empty")
        db.add_channels(["demo"])
        row = db.list_channels()[0]
        self.assertEqual(row["scanned"], 0)
        self.assertIsNone(row["total_views"])
        self.assertIsNone(row["avg_views"])
        self.assertIsNone(row["er"])
        self.assertIsNone(row["d_followers"])

    def test_snapshot_total_views_counts_every_stored_video(self):
        db = _fresh_db("snapshot")
        db.add_channels(["demo"])
        db.finish_scan("demo", [_video("a", play=100), _video("b", play=50)], full=False)
        # Lần quét nhanh chỉ thấy một video nhưng tổng view vẫn là của cả kênh.
        db.finish_scan("demo", [_video("a", play=120)], full=False)
        row = db.list_channels()[0]
        self.assertEqual(row["total_views"], 170)
        self.assertEqual(row["d_views"], 20)

    def test_full_scan_drops_videos_that_disappeared(self):
        db = _fresh_db("full")
        db.add_channels(["demo"])
        db.finish_scan("demo", [_video("a"), _video("gone")], full=False)
        db.finish_scan("demo", [_video("a")], full=False)
        self.assertEqual({v["id"] for v in db.list_videos("demo")}, {"a", "gone"})
        self.assertIsNone(db.get_channel("demo")["full_scan_at"])
        db.finish_scan("demo", [_video("a")], full=True)
        self.assertEqual([v["id"] for v in db.list_videos("demo")], ["a"])
        self.assertIsNotNone(db.get_channel("demo")["full_scan_at"])

    def test_settings_ignore_unknown_keys(self):
        db = _fresh_db("settings")
        db.set_settings({"show_browser": 1, "auto_refresh_minutes": "15", "bogus": "x"})
        got = db.get_settings()
        self.assertEqual((got["show_browser"], got["auto_refresh_minutes"]), ("1", "15"))
        self.assertNotIn("bogus", got)

    def test_interrupted_jobs_become_errors_on_start(self):
        db = _fresh_db("interrupted")
        db.add_channels(["demo"])
        db.set_status("demo", "running")
        db.init()
        channel = db.get_channel("demo")
        self.assertEqual(channel["status"], "error")
        self.assertIn("Bị dừng giữa chừng", channel["message"])

    def test_old_database_gets_new_columns(self):
        import sqlite3

        import db as dbmod

        data = Path(tempfile.mkdtemp(prefix="migrate-", dir=DATA))
        path = data / "data.db"
        old = sqlite3.connect(path)
        legacy_schema = dbmod.SCHEMA.replace(
            ",\n    last_scan_at INTEGER,\n    last_scan_label TEXT DEFAULT ''", ""
        )
        self.assertNotIn("last_scan_at", legacy_schema)
        old.executescript(legacy_schema)
        old.execute("INSERT INTO channels(username) VALUES('legacy')")
        old.commit()
        old.close()
        dbmod.DB_PATH = path
        dbmod.init()
        dbmod.init()
        columns = {row[1] for row in sqlite3.connect(path).execute("PRAGMA table_info(channels)")}
        self.assertTrue({"last_scan_at", "last_scan_label"} <= columns)
        self.assertEqual(dbmod.get_channel("legacy")["username"], "legacy")


class SchemaVersionTest(unittest.TestCase):
    def test_connection_is_closed_so_directory_can_be_removed(self):
        db = _fresh_db("closed")
        db.add_channels(["demo"])
        db.list_channels()
        db.get_settings()
        shutil.rmtree(db.DB_PATH.parent)  # trên Windows sẽ PermissionError nếu còn kết nối mở

    def test_new_and_legacy_databases_reach_the_same_version(self):
        import sqlite3

        db = _fresh_db("version-new")
        with db._conn() as c:
            self.assertEqual(db.schema_version(c), db.SCHEMA_VERSION)
        db.init()
        with db._conn() as c:
            self.assertEqual(db.schema_version(c), db.SCHEMA_VERSION)

        legacy = Path(tempfile.mkdtemp(prefix="version-legacy-", dir=DATA)) / "data.db"
        old = sqlite3.connect(legacy)
        old.executescript(db.SCHEMA.replace(
            ",\n    last_scan_at INTEGER,\n    last_scan_label TEXT DEFAULT ''", ""))
        old.execute("INSERT INTO channels(username) VALUES('legacy')")
        old.commit()
        self.assertEqual(old.execute("PRAGMA user_version").fetchone()[0], 0)
        old.close()
        db.DB_PATH = legacy
        db.init()
        with db._conn() as c:
            self.assertEqual(db.schema_version(c), db.SCHEMA_VERSION)
            tables = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        self.assertIn("alerts", tables)
        self.assertEqual(db.get_channel("legacy")["username"], "legacy")

    def test_migrating_an_old_database_makes_a_backup_first(self):
        import sqlite3

        db = _fresh_db("backup-premigrate")
        with db._conn() as c:
            c.execute("PRAGMA user_version=0")
        db.add_channels(["keep_me"])
        before = len(db.list_backups())
        db.init()
        backups = db.list_backups()
        self.assertEqual(len(backups), before + 1)
        self.assertIn("-v0", backups[0].name)
        copy = sqlite3.connect(backups[0])
        try:
            self.assertEqual(copy.execute("SELECT username FROM channels").fetchall(), [("keep_me",)])
        finally:
            copy.close()

    def test_daily_backup_and_rotation(self):
        import os
        import time

        db = _fresh_db("backup-daily")
        db.add_channels(["demo"])
        db.init()  # có DB và chưa có bản sao lưu nào: tạo một bản
        self.assertEqual(len(db.list_backups()), 1)
        db.init()  # vừa sao lưu xong: không tạo thêm
        self.assertEqual(len(db.list_backups()), 1)
        for i in range(10):
            path = db.backup_now(f"t{i}")
            self.assertIsNotNone(path)
            os.utime(path, (time.time() - 100 + i, time.time() - 100 + i))
        self.assertEqual(len(db.list_backups()), db.BACKUP_KEEP)

    def test_backup_failure_does_not_block_startup(self):
        db = _fresh_db("backup-fail")
        db.add_channels(["demo"])
        blocker = db.DB_PATH.parent / "backups"
        blocker.write_text("not a directory", encoding="utf-8")
        self.assertIsNone(db.backup_now())
        db.init()
        self.assertTrue(db.channel_exists("demo"))


class JobTest(unittest.TestCase):
    def test_normalize_job(self):
        from scraper import normalize_job

        self.assertEqual(normalize_job({"mode": "quick"}), {"mode": "quick"})
        self.assertEqual(
            normalize_job({"mode": "videos", "ids": ["1", 2, "x", "1", ""]}),
            {"mode": "videos", "ids": ["1", "2"]},
        )
        scan = normalize_job({"mode": "scan", "limit": "5", "since": 900, "until": 100})
        self.assertEqual((scan["limit"], scan["since"], scan["until"]), (5, 100, 900))
        by_days = normalize_job({"mode": "scan", "days": 7})
        self.assertEqual(by_days["until"], 0)
        self.assertGreater(by_days["since"], 0)
        self.assertEqual(normalize_job({"mode": "scan", "limit": -3})["limit"], 0)

    def test_describe_job(self):
        from scraper import describe_job, normalize_job

        self.assertEqual(describe_job({"mode": "quick"}), "Cập nhật nhanh")
        self.assertEqual(describe_job(normalize_job({"mode": "scan"})), "Quét toàn bộ video")
        self.assertEqual(describe_job(normalize_job({"mode": "scan", "days": 7})), "Quét 7 ngày qua")
        self.assertEqual(describe_job(normalize_job({"mode": "scan", "limit": 3})), "Quét 3 video gần nhất")
        self.assertEqual(
            describe_job(normalize_job({"mode": "videos", "ids": ["1", "2"]})), "Cập nhật view 2 video"
        )


class QueueTest(unittest.TestCase):
    def setUp(self):
        self.db = _fresh_db("queue")
        self.db.add_channels(["demo", "other"])
        self.scraper = _make_scraper()

    def test_quick_does_not_replace_pending_job_but_scan_upgrades_it(self):
        s = self.scraper
        self.assertEqual(s.enqueue(["demo"], {"mode": "quick"}), 1)
        self.assertEqual(s.enqueue(["demo"], {"mode": "quick"}), 0)
        self.assertEqual(s.enqueue(["demo"], {"mode": "scan", "limit": 5}), 1)
        self.assertEqual(s.pending["demo"]["mode"], "scan")
        self.assertEqual(s.jobs.qsize(), 1)
        self.assertEqual(s.enqueue(["demo"], {"mode": "quick"}), 0)
        self.assertEqual(s.pending["demo"]["mode"], "scan")

    def test_video_jobs_merge_ids_in_their_own_slot(self):
        s = self.scraper
        s.enqueue(["demo"], {"mode": "videos", "ids": ["1", "2"]})
        s.enqueue(["demo"], {"mode": "videos", "ids": ["2", "3"]})
        s.enqueue(["demo"], {"mode": "quick"})
        self.assertEqual(s.pending["demo#videos"], {"mode": "videos", "ids": ["1", "2", "3"]})
        self.assertEqual(s.pending["demo"], {"mode": "quick"})
        self.assertEqual(s.jobs.qsize(), 2)

    def test_status_matches_what_the_page_reads(self):
        s = self.scraper
        self.assertEqual(s.status(), {"queue": 0, "current": None})
        s.enqueue(["demo", "other"], {"mode": "quick"})
        self.assertEqual(s.status()["queue"], 2)
        s.current = {"username": "demo", "label": "Cập nhật nhanh", "videos": 0, "total": 0, "started": 1}
        status = s.status()
        self.assertEqual(status["queue"], 3)
        self.assertEqual(status["current"]["label"], "Cập nhật nhanh")
        self.assertIsNot(status["current"], s.current)

    def test_deleted_channel_is_skipped(self):
        s = self.scraper
        calls = []
        s._scrape = lambda key, job: calls.append(key)
        s.enqueue(["demo", "other"], {"mode": "quick"})
        self.db.delete_channels(["demo"])
        self.assertTrue(_run_worker(s, lambda: s.jobs.empty() and calls))
        self.assertEqual(calls, ["other"])

    def test_videos_mode_skips_channel_scrape(self):
        s = self.scraper
        calls = []
        s._update_videos = lambda key, ids: calls.append(("videos", key, ids))
        s._scrape = lambda key, job: calls.append(("scrape", key))
        s.enqueue(["demo"], {"mode": "videos", "ids": ["1", "2"]})
        self.assertTrue(_run_worker(s, lambda: calls))
        self.assertEqual(calls, [("videos", "demo", ["1", "2"])])

    def test_job_queued_while_running_goes_back_to_waiting(self):
        s = self.scraper
        seen = []

        def scrape(key, job):
            if seen:
                return
            s.enqueue([key], {"mode": "scan", "limit": 1})
            seen.append(self.db.get_channel(key)["status"])

        s._scrape = scrape
        s.enqueue(["demo"], {"mode": "quick"})
        # Dừng worker trước khi nó nhặt job vừa xếp: chỉ cần xem trạng thái ngay sau job đầu.
        gate = threading.Event()
        real_set_status = self.db.set_status

        def set_status(key, status, *rest):
            real_set_status(key, status, *rest)
            if status == "queued" and seen:
                gate.set()

        with mock.patch.object(self.db, "set_status", set_status):
            with mock.patch("time.sleep", lambda seconds: gate.wait(5) and threading.Event().wait(60)):
                threading.Thread(target=s._run, daemon=True).start()
                self.assertTrue(gate.wait(10))
        self.assertEqual(seen, ["running"])
        self.assertEqual(self.db.get_channel("demo")["status"], "queued")
        self.assertEqual(s.pending["demo"]["mode"], "scan")


class _Response:
    def __init__(self, url, payload):
        self.url = url
        self._payload = payload

    def json(self):
        return self._payload


class _Mouse:
    def __init__(self):
        self.wheels = []

    def wheel(self, x, y):
        self.wheels.append((x, y))


class _Locator:
    def __init__(self, selector, html):
        self.selector = selector
        self.html = html

    def count(self):
        if "#captcha-container" in self.selector and 'id="captcha-container"' in self.html:
            return 1
        if "captcha_verify_container" in self.selector and "captcha_verify_container" in self.html:
            return 1
        if "[class*='captcha']" in self.selector and 'class="secsdk-captcha"' in self.html:
            return 1
        return 0


class _Page:
    html = "<html>captcha word only</html>"

    def __init__(self, payloads):
        self._payloads = list(payloads)
        self.mouse = _Mouse()
        self._handler = None
        self.scrolls = 0
        self.closed = False

    def on(self, event, handler):
        self._handler = handler

    def goto(self, url, **kwargs):
        self._emit()

    def evaluate(self, script):
        if "scrollTo" in script:
            self.scrolls += 1
            self._emit()
            return None
        if "UNIVERSAL_DATA" in script:
            return (
                '{"__DEFAULT_SCOPE__":{"webapp.user-detail":{"userInfo":{"user":'
                '{"uniqueId":"demo","nickname":"Demo","privateAccount":false,"secUid":"s",'
                '"verified":false,"signature":"","createTime":1},"stats":{"followerCount":3,'
                '"followingCount":1,"heartCount":2,"videoCount":2}}}}}'
            )
        return None

    def wait_for_timeout(self, ms):
        return None

    def close(self):
        self.closed = True

    def locator(self, selector):
        return _Locator(selector, self.html)

    def _emit(self):
        if self._payloads and self._handler:
            self._handler(self._payloads.pop(0))


class _Context:
    def __init__(self, page):
        self.page = page

    def new_page(self):
        return self.page


def _item_list(*items, has_more=False):
    return _Response(
        "https://www.tiktok.com/api/post/item_list?cursor=0",
        {"hasMore": has_more, "itemList": [
            {"id": i[0], "desc": "", "createTime": i[1], "stats": {"playCount": i[2]}, "video": {}} for i in items
        ]},
    )


def _scrape(name, page, job, setup=None):
    from scraper import normalize_job

    db = _fresh_db(name)
    db.add_channels(["demo"])
    if setup:
        setup(db)
    scraper = _make_scraper()
    scraper._browser = lambda: _Context(page)
    scraper._scrape("demo", normalize_job(job))
    return db, scraper


class ScrapeTest(unittest.TestCase):
    def test_scan_stops_when_has_more_is_false(self):
        page = _Page([_item_list(("2", 200, 5), ("1", 100, 9))])
        db, _ = _scrape("scroll", page, {"mode": "scan"})
        self.assertEqual(page.scrolls, 0)
        self.assertEqual(page.mouse.wheels, [])
        self.assertEqual([v["id"] for v in db.list_videos("demo")], ["2", "1"])
        channel = db.get_channel("demo")
        self.assertEqual(channel["status"], "done")
        self.assertEqual(channel["followers"], 3)
        self.assertTrue(page.closed)

    def test_scan_keeps_scrolling_while_tiktok_has_more(self):
        page = _Page([
            _item_list(("3", 300, 1), has_more=True),
            _item_list(("2", 200, 1), has_more=True),
            _item_list(("1", 100, 1), has_more=False),
        ])
        db, _ = _scrape("scroll-more", page, {"mode": "scan"})
        self.assertEqual(page.scrolls, 2)
        self.assertEqual([v["id"] for v in db.list_videos("demo")], ["3", "2", "1"])

    def test_quick_mode_never_scrolls(self):
        page = _Page([_item_list(("2", 200, 1), has_more=True), _item_list(("1", 100, 1))])
        db, _ = _scrape("quick", page, {"mode": "quick"})
        self.assertEqual(page.scrolls, 0)
        self.assertEqual([v["id"] for v in db.list_videos("demo")], ["2"])
        self.assertEqual(db.get_channel("demo")["message"], "")

    def test_limit_counts_pinned_videos_extra(self):
        # Ba video ghim cũ nằm đầu danh sách: dừng ở limit + 3 để vẫn thấy video mới nhất.
        pinned = [("p1", 10, 1), ("p2", 11, 1), ("p3", 12, 1)]
        page = _Page([
            _item_list(*pinned, ("n2", 500, 1), has_more=True),
            _item_list(("n1", 600, 1), has_more=True),
            _item_list(("older", 50, 1), has_more=True),
        ])
        db, _ = _scrape("pinned", page, {"mode": "scan", "limit": 2})
        ids = [v["id"] for v in db.list_videos("demo")]
        self.assertEqual(ids, ["n1", "n2"])
        self.assertEqual(page.scrolls, 1)

    def test_since_drops_older_videos_and_pinned_video_does_not_stop_early(self):
        page = _Page([
            _item_list(("pin", 10, 1), ("a", 900, 1), ("b", 800, 1), has_more=True),
            _item_list(("c", 700, 1), ("d", 100, 1), has_more=True),
            _item_list(("e", 50, 1), has_more=False),
        ])
        db, _ = _scrape("since", page, {"mode": "scan", "since": 600})
        self.assertEqual([v["id"] for v in db.list_videos("demo")], ["a", "b", "c"])
        self.assertIsNone(db.get_channel("demo")["full_scan_at"])

    def test_full_scan_removes_vanished_videos_only_when_complete(self):
        def seed(db):
            db.finish_scan("demo", [_video("gone")], full=False)

        db, _ = _scrape("full-ok", _Page([_item_list(("1", 100, 1))]), {"mode": "scan"}, seed)
        self.assertEqual([v["id"] for v in db.list_videos("demo")], ["1"])
        self.assertIsNotNone(db.get_channel("demo")["full_scan_at"])
        # Quét dở (còn has_more nhưng TikTok ngừng trả thêm): không được xóa video cũ.
        pages = [_item_list(("1", 100, 1), has_more=True)]
        db, _ = _scrape("full-partial", _Page(pages), {"mode": "scan"}, seed)
        self.assertEqual({v["id"] for v in db.list_videos("demo")}, {"1", "gone"})
        self.assertIn("TikTok ngừng trả thêm", db.get_channel("demo")["message"])

    def test_private_channel_has_no_videos(self):
        class Private(_Page):
            def evaluate(self, script):
                return (
                    '{"__DEFAULT_SCOPE__":{"webapp.user-detail":{"userInfo":{"user":'
                    '{"uniqueId":"demo","privateAccount":true},"stats":{"videoCount":4}}}}}'
                )

        db, _ = _scrape("private", Private([]), {"mode": "scan"})
        channel = db.get_channel("demo")
        self.assertEqual(channel["status"], "done")
        self.assertEqual(channel["private"], 1)
        self.assertIn("riêng tư", channel["message"])

    def test_missing_channel_reports_reason(self):
        class Missing(_Page):
            def evaluate(self, script):
                return '{"__DEFAULT_SCOPE__":{"webapp.user-detail":{"statusCode":10221}}}'

        db, _ = _scrape("missing", Missing([]), {"mode": "quick"})
        channel = db.get_channel("demo")
        self.assertEqual(channel["status"], "error")
        self.assertEqual(channel["message"], "Kênh không tồn tại hoặc bị khóa")

    def test_progress_is_visible_while_scanning(self):
        page = _Page([_item_list(("2", 200, 5), ("1", 100, 9))])

        from scraper import normalize_job

        db = _fresh_db("progress")
        db.add_channels(["demo"])
        scraper = _make_scraper()
        scraper.current = {"username": "demo", "label": "", "videos": 0, "total": 0, "started": 0}
        scraper._browser = lambda: _Context(page)
        scraper._scrape("demo", normalize_job({"mode": "scan"}))
        self.assertEqual(scraper.current["videos"], 2)


class _NavigatingPage(_Page):
    def __init__(self, payloads):
        super().__init__(payloads)
        self.navigations = 0
        self.load_waits = 0

    def evaluate(self, script):
        if "UNIVERSAL_DATA" in script and not self.navigations:
            self.navigations += 1
            raise RuntimeError(
                "Page.evaluate: Execution context was destroyed, most likely because of a navigation"
            )
        return super().evaluate(script)

    def wait_for_load_state(self, state, timeout=None):
        self.load_waits += 1


class NavigationTest(unittest.TestCase):
    def test_scan_survives_navigation_during_evaluate(self):
        page = _NavigatingPage([_item_list(("1", 100, 9))])
        db, _ = _scrape("navigation", page, {"mode": "scan"})
        self.assertEqual(page.navigations, 1)
        self.assertEqual(page.load_waits, 1)
        self.assertEqual([v["id"] for v in db.list_videos("demo")], ["1"])
        self.assertEqual(db.get_channel("demo")["status"], "done")

    def test_other_evaluate_errors_still_raise(self):
        from scraper import _evaluate_after_navigation

        class Broken:
            def evaluate(self, script):
                raise RuntimeError("Target page, context or browser has been closed")

        with self.assertRaises(RuntimeError):
            _evaluate_after_navigation(Broken(), "() => 1")


class _DetailPage:
    def __init__(self):
        self.urls = []
        self.closed = False

    def goto(self, url, **kwargs):
        self.urls.append(url)

    def evaluate(self, script):
        if "UNIVERSAL_DATA" not in script:
            return None
        video_id = self.urls[-1].rsplit("/", 1)[-1]
        if video_id == "missing":
            return None
        return (
            '{"__DEFAULT_SCOPE__":{"webapp.video-detail":{"itemInfo":{"itemStruct":'
            '{"id":"%s","desc":"kept","createTime":300,"stats":{"playCount":7},"video":{}}}}}}'
            % video_id
        )

    def wait_for_timeout(self, ms):
        return None

    def close(self):
        self.closed = True


class UpdateVideosTest(unittest.TestCase):
    def test_updates_found_video_and_counts_missing(self):
        db = _fresh_db("update")
        db.add_channels(["demo"])
        page = _DetailPage()
        scraper = _make_scraper()
        scraper.current = {"username": "demo", "videos": 0, "total": 2}
        scraper._browser = lambda: type("Ctx", (), {"new_page": lambda self: page})()
        scraper._update_videos("demo", ["kept", "missing"])
        self.assertTrue(page.closed)
        self.assertEqual(scraper.current["videos"], 2)
        self.assertEqual([video["id"] for video in db.list_videos("demo")], ["kept"])
        self.assertTrue(page.urls[0].endswith("/@demo/video/kept"))
        self.assertTrue(page.urls[1].endswith("/@demo/video/missing"))
        channel = db.get_channel("demo")
        self.assertEqual(channel["status"], "done")
        self.assertIn("Cập nhật view 1 video", channel["message"])
        self.assertIn("1 video không lấy được", channel["message"])


class CaptchaTest(unittest.TestCase):
    def test_word_in_html_is_not_a_widget(self):
        from scraper import Scraper

        page = _DetailPage()
        page.html = "<html>captcha word only</html>"
        page.locator = lambda selector: _Locator(selector, page.html)
        self.assertFalse(Scraper._has_captcha(page))
        page.html = '<div id="captcha-container"></div>'
        self.assertTrue(Scraper._has_captcha(page))

    def test_captcha_marks_channel_and_queue_moves_on(self):
        class CaptchaPage(_Page):
            html = '<div id="captcha-container"></div>'

        db = _fresh_db("captcha-queue")
        db.add_channels(["blocked", "ok"])
        scraper = _make_scraper()
        pages = {"blocked": CaptchaPage([]), "ok": _Page([_item_list(("9", 1, 1))])}
        order = []
        real_scrape = scraper._scrape

        def scrape(key, job):
            order.append(key)
            return real_scrape(key, job)

        scraper._scrape = scrape
        scraper._browser = lambda: _Context(pages[order[-1]])
        scraper.enqueue(["blocked", "ok"], {"mode": "quick"})

        def finished():
            return all(db.get_channel(k)["status"] in ("done", "error") for k in ("blocked", "ok"))

        self.assertTrue(_run_worker(scraper, finished))
        self.assertEqual(order, ["blocked", "ok"])
        blocked = db.get_channel("blocked")
        self.assertEqual(blocked["status"], "error")
        self.assertIn("captcha", blocked["message"])
        self.assertNotIn("Lỗi:", blocked["message"])
        self.assertIsNone(blocked["last_scan_at"])
        self.assertEqual(db.get_channel("ok")["status"], "done")


class RetryTest(unittest.TestCase):
    def _run_one(self, name, failures):
        db = _fresh_db(name)
        db.add_channels(["demo"])
        scraper = _make_scraper()
        calls = []

        def scrape(key, job):
            calls.append(key)
            if len(calls) <= len(failures):
                raise failures[len(calls) - 1]
            db.set_status(key, "done", "")

        scraper._scrape = scrape
        scraper.enqueue(["demo"], {"mode": "scan", "limit": 3})
        self.assertTrue(_run_worker(scraper, lambda: db.get_channel("demo")["status"] in ("done", "error")))
        return calls, db.get_channel("demo")

    def test_transient_error_is_retried_then_succeeds(self):
        calls, channel = self._run_one("retry-ok", [RuntimeError("Timeout 30000ms exceeded")])
        self.assertEqual(len(calls), 2)
        self.assertEqual(channel["status"], "done")
        self.assertEqual(channel["last_scan_label"], "Quét 3 video gần nhất")
        self.assertIsNotNone(channel["last_scan_at"])

    def test_transient_error_gives_up_after_three_attempts(self):
        errors = [RuntimeError("net::ERR_CONNECTION_RESET")] * 3
        calls, channel = self._run_one("retry-fail", errors)
        self.assertEqual(len(calls), 3)
        self.assertEqual(channel["status"], "error")
        self.assertIn("đã thử 3 lần", channel["message"])

    def test_permanent_and_profile_errors_are_not_retried(self):
        from scraper import ProfileBusyError

        for name, error in (
            ("no-retry-missing", RuntimeError("Kênh không tồn tại")),
            ("no-retry-profile", ProfileBusyError("Hồ sơ trình duyệt đang bị giữ")),
        ):
            calls, channel = self._run_one(name, [error])
            self.assertEqual(len(calls), 1, name)
            self.assertEqual(channel["status"], "error", name)
            self.assertNotIn("đã thử", channel["message"], name)


class InstanceLockTest(unittest.TestCase):
    def test_second_app_on_same_data_dir_is_refused(self):
        import subprocess

        import instance_lock

        data = Path(tempfile.mkdtemp(prefix="lock-", dir=DATA))
        held = instance_lock.acquire(data)
        try:
            probe = (
                "import sys, instance_lock\n"
                "try:\n"
                "    instance_lock.acquire(sys.argv[1])\n"
                "except instance_lock.AlreadyRunning as exc:\n"
                "    print('busy', exc)\n"
                "else:\n"
                "    print('acquired')\n"
            )
            out = subprocess.run(
                [sys.executable, "-c", probe, str(data)],
                cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=30,
            ).stdout
            self.assertTrue(out.startswith("busy"), out)
            self.assertIn(f"PID {os.getpid()}", out)
        finally:
            held.release()
        again = instance_lock.acquire(data)
        again.release()
        self.assertFalse((data / instance_lock.PID_NAME).exists())


class ApiTest(unittest.TestCase):
    def _client(self, name):
        import app as appmod
        from fastapi.testclient import TestClient

        db = _fresh_db(name)
        scraper = _make_scraper()
        appmod.scraper = scraper
        return TestClient(appmod.app), db, scraper, appmod

    def test_ping_identifies_the_tool(self):
        client, *_ = self._client("api-ping")
        self.assertEqual(client.get("/api/ping").json(), {"app": "theo-doi-kenh-tiktok"})

    def test_add_list_and_refresh_channels(self):
        client, db, scraper, _ = self._client("api-channels")
        res = client.post("/api/channels", json={"text": "@Demo bad! https://www.tiktok.com/@Other",
                                                 "scan": {"mode": "quick"}})
        self.assertEqual(res.json(), {"added": ["demo", "other"], "existed": [], "invalid": ["bad!"]})
        self.assertEqual(scraper.jobs.qsize(), 2)
        listing = client.get("/api/channels").json()
        self.assertEqual([c["username"] for c in listing["channels"]], ["demo", "other"])
        self.assertEqual(listing["worker"], {"queue": 2, "current": None})
        # Kênh không có trong danh sách bị bỏ qua; job là bắt buộc.
        res = client.post("/api/channels/refresh", json={"usernames": ["Demo", "ghost"], "job": {"mode": "scan"}})
        self.assertEqual(res.json(), {"queued": 1})
        self.assertEqual(client.post("/api/channels/refresh", json={"usernames": ["demo"]}).status_code, 422)

    def test_meta_delete_and_videos(self):
        client, db, scraper, _ = self._client("api-meta")
        db.add_channels(["demo"])
        client.post("/api/channels/meta", json={"usernames": ["DEMO"], "group_name": "A", "note": "n"})
        channel = client.get("/api/channels/demo/videos").json()["channel"]
        self.assertEqual((channel["group_name"], channel["note"]), ("A", "n"))
        self.assertEqual(client.get("/api/channels/ghost/videos").status_code, 404)
        self.assertEqual(client.post("/api/channels/delete", json={"usernames": ["demo"]}).json(), {"ok": True})
        self.assertIsNone(db.get_channel("demo"))

    def test_open_only_allows_tiktok_and_the_tool_itself(self):
        client, *_ = self._client("api-open")
        opened = []
        urls = [
            "https://www.tiktok.com/@a/video/1",
            "https://evil.example/",
            "file:///C:/Windows/System32/calc.exe",
            "http://127.0.0.1:8765/videos?u=a",
        ]
        with mock.patch("webbrowser.open_new_tab", opened.append):
            client.post("/api/open", json={"urls": urls})
        self.assertEqual(opened, [urls[0], urls[3]])

    def test_settings_roundtrip(self):
        client, *_ = self._client("api-settings")
        res = client.post("/api/settings", json={"show_browser": "1", "auto_refresh_minutes": 30})
        got = res.json()
        self.assertEqual((got["show_browser"], got["auto_refresh_minutes"]), ("1", "30"))
        self.assertEqual(client.get("/api/settings").json(), res.json())


class ExportTest(unittest.TestCase):
    def _client(self, name):
        import app as appmod
        from fastapi.testclient import TestClient

        db = _fresh_db(name)
        appmod.scraper = _make_scraper()
        return TestClient(appmod.app), db

    def test_channels_xlsx_has_vietnamese_columns(self):
        import io
        import re

        from openpyxl import load_workbook

        client, db = self._client("export-channels")
        db.add_channels(["demo", "other"])
        db.save_channel_info("demo", {"nickname": "Demo", "followers": 7, "verified": 1})
        db.finish_scan("demo", [_video("v1", play=10, digg=1)], full=False)
        res = client.get("/api/export/channels.xlsx?usernames=demo")
        self.assertEqual(res.status_code, 200)
        self.assertRegex(res.headers["content-disposition"], r'filename="kenh_tiktok_\d{8}_\d{4}\.xlsx"')
        sheet = load_workbook(io.BytesIO(res.content)).active
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(len(rows), 2)
        head = rows[0]
        self.assertEqual(len(head), 25)
        self.assertEqual(head[:4], ["Username", "Tên kênh", "Link", "Follow"])
        row = dict(zip(head, rows[1]))
        self.assertEqual(row["Link"], "https://www.tiktok.com/@demo")
        self.assertEqual(row["Follow"], 7)
        self.assertEqual(row["Tick xanh"], "Có")
        self.assertIn(row["Riêng tư"], (None, ""))
        self.assertEqual(row["Video đã quét"], 1)
        self.assertEqual(row["Tỷ lệ tương tác %"], 10.0)
        self.assertTrue(re.match(r"\d\d/\d\d/\d{4} \d\d:\d\d", row["Cập nhật lúc"]))

    def test_csv_matches_xlsx_and_neutralises_formulas(self):
        import csv
        import io

        from openpyxl import load_workbook

        client, db = self._client("export-csv")
        db.add_channels(["demo"])
        videos = [
            _video("v1", 1_700_000_001, 100, description="=SUM(A1)", music="@mix"),
            _video("v2", 1_700_000_002, -7, description='Xin chào, "Việt"\nNam', music="-x"),
            _video("v3", 1_700_000_003, 0, description="+1 ok", music=""),
        ]
        db.finish_scan("demo", videos, full=False)

        res = client.get("/api/export/videos/demo.csv")
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.headers["content-type"].startswith("text/csv"))
        self.assertRegex(res.headers["content-disposition"], r'filename="video_demo_\d{8}_\d{4}\.csv"')
        self.assertTrue(res.content.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(io.StringIO(res.content.decode("utf-8-sig"), newline="")))

        sheet = load_workbook(io.BytesIO(client.get("/api/export/videos/demo.xlsx").content)).active
        xlsx = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(rows[0], xlsx[0])
        self.assertEqual(rows[0][:2], ["Link", "Mô tả"])
        self.assertEqual([row[0].rsplit("/", 1)[-1] for row in rows[1:]], ["v3", "v2", "v1"])
        by_id = {row[0].rsplit("/", 1)[-1]: row for row in rows[1:]}
        self.assertEqual(by_id["v1"][1], "'=SUM(A1)")
        self.assertEqual(by_id["v1"][11], "'@mix")
        self.assertEqual(by_id["v2"][1], 'Xin chào, "Việt"\nNam')
        self.assertEqual(by_id["v2"][11], "'-x")
        self.assertEqual(by_id["v2"][4], "-7")
        self.assertEqual(by_id["v3"][1], "'+1 ok")
        self.assertEqual(by_id["v3"][11], "")

    def test_unknown_channel_is_404(self):
        client, _ = self._client("export-404")
        self.assertEqual(client.get("/api/export/videos/nobody.csv").status_code, 404)
        self.assertEqual(client.get("/api/export/videos/nobody.xlsx").status_code, 404)


if __name__ == "__main__":
    unittest.main()
